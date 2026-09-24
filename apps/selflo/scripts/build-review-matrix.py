#!/usr/bin/env python3
import argparse, json, re
from pathlib import Path
from openpyxl import load_workbook

CURATED = {
    'research.dhammapada.day11.v019': (94, 'Hình ảnh mạnh; phân biệt biết và sống; source verse rõ.'),
    'research.dhammapada.day11.v064': (92, 'Độc lập tốt; dễ nhớ; chiều sâu về trải nghiệm và tri thức.'),
    'research.mn.day12.mn22_02': (90, 'Perspective sâu về công cụ và sự buông bỏ; cần kiểm tra overlap story.'),
    'research.laozi.ch12.034': (89, 'Rất hợp attention overload; nguyên văn rõ; đọc độc lập tốt.'),
    'research.dhammapada.day11.v276': (88, 'Khớp triết lý Selflo: người khác chỉ đường nhưng không bước thay.'),
    'research.laozi.ch02.007': (87, 'Góc nhìn mạnh về hành động không chiếm hữu; primary text rõ.'),
    'research.laozi.ch13.037': (86, 'Ngắn, độc lập và liên quan trực tiếp tới đánh giá bên ngoài.'),
    'research.laozi.ch17.054': (85, 'Perspective khác biệt về leadership và agency; cần giữ ngữ cảnh chương.'),
    'research.laozi.ch20.061': (84, 'Có chiều sâu về bất định và cảm giác không theo kịp số đông.'),
    'research.dhammapada.day11.v043': (83, 'Source verse rõ; cần review bản chuyển ngữ Selflo.'),
    'research.dhammapada.day11.v223': (82, 'Hữu ích cho giận dữ và agency; wording là chuyển ý một phần kệ.'),
    'research.mn.day12.mn20_01': (81, 'Perspective thực tế về chuyển hướng chú ý; tránh diễn đạt như lời khuyên lâm sàng.'),
    'research.laozi.ch05.014': (80, 'Ngắn và hợp attention; bản Việt hiện mang tính diễn giải.'),
}

def clean(v):
    return '' if v is None else str(v).strip()

def sheet_records(wb, name):
    rows = wb[name].iter_rows(values_only=True)
    headers = [clean(v) for v in next(rows)]
    out = []
    for values in rows:
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        if any(v is not None and clean(v) for v in values): out.append(record)
    return out

def load_quotes(channel_dir):
    manifest = json.loads((channel_dir/'manifest.json').read_text())
    quotes = []
    for f in manifest.get('files', []):
        if f.get('kind') != 'quote_pack': continue
        doc = json.loads((channel_dir/f['path']).read_text())
        quotes.extend(doc.get('quotes', []))
    return manifest, quotes

def confidence(value):
    v = clean(value).lower()
    if v.startswith('a') or 'verified' in v or v in {'public_domain','selflo_owned','licensed','permission_granted'}: return 'high'
    if v.startswith('b') or 'likely' in v: return 'medium'
    return 'low'

def research_score(q, n, m):
    ident = clean(q.get('Quote ID'))
    if ident in CURATED: return CURATED[ident]
    score = {'high': 68, 'medium': 56, 'low': 42}[confidence(q.get('Mức xác minh'))]
    fit = clean(q.get('Mức phù hợp Selflo')).lower()
    score += 12 if fit == 'high' else 6 if fit == 'medium' else 0
    if clean(q.get('Source URL')): score += 4
    if clean(q.get('Release readiness')).lower() in {'ready','release ready','yes'}: score += 5
    if clean(n.get('Exact duplicate cluster ID')): score -= 15
    if clean(n.get('Near-duplicate cluster ID')): score -= 8
    text = clean(q.get('Bản dịch / bản làm việc tiếng Việt'))
    if re.search(r'đoạn|văn bản|theo đoạn này|hai mặt ấy', text, re.I): score -= 8
    why = 'Điểm sàng lọc từ source, Selflo fit, readiness và cảnh báo duplicate; cần editorial review trước khi chốt.'
    return max(0, min(99, score)), why

def exclusion_reason(q, n):
    if clean(q.get('Mức phù hợp Selflo')) != 'High':
        return 'selflo_fit_not_high', 'Chưa phù hợp Selflo ở mức High.'
    source_url = clean(q.get('Source URL'))
    if not source_url:
        return 'missing_source_url', 'Thiếu đường dẫn nguồn.'
    if not source_url.lower().startswith(('http://', 'https://')):
        return 'source_url_not_public_http', 'Nguồn không phải đường dẫn công khai có thể kiểm tra.'
    if not clean(q.get('Mức xác minh')).startswith(('A', 'B')):
        return 'source_below_B', 'Mức xác minh nguồn thấp hơn B.'
    if clean(q.get('Rủi ro giáo điều / directive')).lower().startswith(('high', 'cao')):
        return 'directive_risk_high', 'Rủi ro diễn đạt giáo điều hoặc ra lệnh cao.'
    text = clean(q.get('Bản dịch / bản làm việc tiếng Việt'))
    if not 15 <= len(text) <= 600:
        return 'text_length_outside_gate', 'Độ dài chưa phù hợp để biên tập thành quote.'
    if clean(n.get('Exact duplicate cluster ID')):
        return 'exact_duplicate_alternate', 'Biến thể trùng hoàn toàn; đang giữ lại để đối chiếu, không nhập Authoring.'
    return 'awaiting_authoring_review', 'Chưa được review vòng Canonical → Authoring.'

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--root', type=Path, default=Path(__file__).resolve().parent.parent)
    ap.add_argument('--workbook', type=Path)
    args = ap.parse_args(); root = args.root.resolve()
    book = args.workbook or root/'quote-chatpgt/Selflo_Content_Master_Post90_Phase2_5.xlsx'
    wb = load_workbook(book, read_only=True, data_only=True)
    raw = sheet_records(wb, 'Quote Library v2')
    normalized = {clean(x.get('Quote ID')): x for x in sheet_records(wb, 'Normalized Corpus')}
    mapping = {clean(x.get('Research Quote ID')): x for x in sheet_records(wb, 'Research Repo Mapping')}
    ledger_path=root/'quote-research/review-pipeline/canonical-ai-review.json'
    ledger=json.loads(ledger_path.read_text()) if ledger_path.exists() else {'decisions':[]}
    ai_reviews={x['quote_id']:x for x in ledger.get('decisions',[])}
    authoring_manifest, authoring = load_quotes(root/'perspective-library/authoring/vi')
    release_manifest, release = load_quotes(root/'perspective-library/release/vi')
    release_ids = {clean(x.get('id')) for x in release}
    authoring_by_id = {clean(x.get('id')): x for x in authoring}
    research_by_id={clean(q.get('Quote ID')):q for q in raw}
    canonical_to_research={}
    for q in raw:
        ident=clean(q.get('Quote ID')); m=mapping.get(ident,{})
        mapped_id=clean(m.get('Canonical ID'))
        canonical_id=mapped_id or (ident if ident in authoring_by_id else '')
        if canonical_id in authoring_by_id:
            canonical_to_research.setdefault(canonical_id, ident)

    rows=[]
    for q in authoring:
        ident=clean(q.get('id')); a=q.get('authorship') or {}; rights=q.get('rights') or {}; review=q.get('review') or {}
        research_id=canonical_to_research.get(ident); rq=research_by_id.get(research_id,{}) if research_id else {}; n=normalized.get(research_id,{}) if research_id else {}
        ai=ai_reviews.get(ident,{}) if ident not in release_ids else {}; ai_score={'ready_for_owner_review':92,'story_review_required':84,'source_verification_required':80,'edit_required':70,'round1_not_pass':30,'source_rights_blocked':20}.get(ai.get('decision'),78)
        rows.append({'record_type':'quote','id':ident,'research_id':research_id,'canonical_id':ident,'text_vi':clean(q.get('text_vi')),'author':clean(a.get('author_name')),'work':clean(a.get('work')),'source_url':clean(a.get('source_url')),'theme':clean(q.get('primary_theme')),'human_experience':clean(n.get('Human experience canonical') or rq.get('Trải nghiệm con người (Human Experience)')),'content_nature':clean(q.get('kind')),'confidence':confidence(rights.get('status')),'verification':clean(a.get('source_detail')),'rights':clean(rights.get('status')),'selflo_fit':clean(rq.get('Mức phù hợp Selflo')),'release_readiness':'Released' if ident in release_ids else 'Needs owner review' if review.get('status')=='needs_owner_review' else clean(review.get('status')),'pipeline':'release' if ident in release_ids else 'authoring','review':clean(review.get('status')),'ai_review_status':ai.get('decision'),'ai_review_round':ai.get('review_round'),'ai_review_note':ai.get('note_vi'),'story_id':q.get('story_id'),'exact_duplicate':clean(n.get('Exact duplicate cluster ID')),'near_duplicate':clean(n.get('Near-duplicate cluster ID')),'issue_code':None,'score':100 if ident in release_ids else ai_score,'why':'Đã phát hành.' if ident in release_ids else ai.get('note_vi') or 'Đang ở Authoring; cần bạn review vòng 2 trước khi Release.'})

    matched_research=set(canonical_to_research.values())
    for q in raw:
        ident=clean(q.get('Quote ID'))
        if ident in matched_research: continue
        n=normalized.get(ident,{}); m=mapping.get(ident,{})
        issue_code,why=exclusion_reason(q,n)
        pipeline='canonical' if issue_code=='awaiting_authoring_review' else 'excluded'
        score,_=research_score(q,n,m)
        rows.append({'record_type':'quote','id':ident,'research_id':ident,'canonical_id':None,'text_vi':clean(q.get('Bản dịch / bản làm việc tiếng Việt')),'author':clean(n.get('Author canonical') or q.get('Tác giả / Attribution')),'work':clean(n.get('Source/work canonical') or q.get('Tên nguồn / tác phẩm')),'source_url':clean(q.get('Source URL')),'theme':clean(n.get('Theme canonical') or q.get('Chủ đề chính (Theme)')),'human_experience':clean(n.get('Human experience canonical') or q.get('Trải nghiệm con người (Human Experience)')),'content_nature':clean(n.get('content_nature canonical') or q.get('Hình thức nội dung (Content Form)')),'confidence':confidence(q.get('Mức xác minh')),'verification':clean(q.get('Mức xác minh')),'rights':clean(q.get('Quyền sử dụng')),'selflo_fit':clean(q.get('Mức phù hợp Selflo')),'release_readiness':clean(q.get('Release readiness')),'pipeline':pipeline,'review':clean(q.get('Owner review')),'ai_review_status':None,'story_id':clean(q.get('Story ID liên quan')) or None,'exact_duplicate':clean(n.get('Exact duplicate cluster ID')),'near_duplicate':clean(n.get('Near-duplicate cluster ID')),'issue_code':issue_code,'score':score,'why':why})
    rows.sort(key=lambda x:(-x['score'],x['record_type'],x['id']))
    for i,x in enumerate(rows,1): x['rank']=i
    stage_counts={stage:sum(1 for x in rows if x['pipeline']==stage) for stage in ('canonical','excluded','authoring','release')}
    summary={'total_quotes':len(rows),'canonical_pending_authoring':stage_counts['canonical'],'excluded_not_authoring':stage_counts['excluded'],'authoring_pending_release':stage_counts['authoring'],'release_quotes':stage_counts['release'],'stage_sum':sum(stage_counts.values()),'research_source_rows':len(raw),'canonical_authoring_quotes':len(authoring),'authoring_revision':authoring_manifest.get('library_revision'),'release_revision':release_manifest.get('library_revision')}
    out={'schema_version':'selflo.review-matrix.v3','generated_from':book.name,'summary':summary,'rows':rows}
    target=root/'review-matrix/data.json';target.parent.mkdir(parents=True,exist_ok=True);target.write_text(json.dumps(out,ensure_ascii=False,separators=(',',':'))+'\n')
    print(json.dumps(summary,ensure_ascii=False,indent=2))

if __name__=='__main__': main()
