#!/usr/bin/env python3
"""Select source-backed Selflo-fit research rows for Authoring review."""
import argparse, datetime as dt, json, re, unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import quote, urlsplit, urlunsplit

from openpyxl import load_workbook


def read(path): return json.loads(path.read_text(encoding='utf-8'))
def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
def clean(value): return str(value or '').strip()
def normalized_text(value): return re.sub(r'\W+', '', clean(value).lower(), flags=re.UNICODE)
def stable_id(value):
    folded=unicodedata.normalize('NFKD', value).encode('ascii','ignore').decode('ascii').lower()
    return re.sub(r'[^a-z0-9]+','_',folded).strip('_').replace('_','.',1) if not re.fullmatch(r'[a-z0-9]+([._][a-z0-9]+)*',value) else value
def records(workbook, sheet):
    rows=workbook[sheet].iter_rows(values_only=True); headers=next(rows)
    return [{headers[i]:value for i,value in enumerate(row) if i < len(headers) and headers[i]} for row in rows]
def uri(value):
    part=urlsplit(clean(value))
    return urlunsplit((part.scheme, part.netloc, quote(part.path, safe='/%'), quote(part.query, safe='=&;%:+,/?@'), quote(part.fragment, safe='')))


def current_quotes(root):
    source=root/'perspective-library/source/vi'; index=read(source/'source.json'); result={}
    for entry in index['files']:
        if entry['kind'] not in ('quote_pack','quote_fragment'): continue
        for item in read(source/entry['path']).get('quotes',[]): result[item['id']]=item
    return result


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument('--workbook', type=Path, default=Path('quote-chatpgt/Selflo_Content_Master_Post90_Phase2_5.xlsx'))
    parser.add_argument('--ledger', type=Path, default=Path('/tmp/selflo-full-research-ledger.json'))
    parser.add_argument('--write', action='store_true')
    args=parser.parse_args(); root=args.root.resolve(); workbook=(root/args.workbook).resolve() if not args.workbook.is_absolute() else args.workbook
    wb=load_workbook(workbook, read_only=True, data_only=True)
    raw=records(wb,'Quote Library v2'); normalized={clean(x.get('Quote ID')):x for x in records(wb,'Normalized Corpus')}
    current=current_quotes(root); current_ids=set(current); current_text={normalized_text(x.get('text_vi')) for x in current.values()}
    reasons=Counter(); preliminary=[]
    for q in raw:
        ident=clean(q.get('Quote ID')); n=normalized.get(ident,{})
        if ident in current_ids: reasons['already_in_canonical']+=1; continue
        if clean(q.get('Mức phù hợp Selflo'))!='High': reasons['selflo_fit_not_high']+=1; continue
        source_url=clean(q.get('Source URL')); source_scheme=urlsplit(source_url).scheme.lower()
        if not source_url: reasons['missing_source_url']+=1; continue
        if source_scheme not in ('http','https'): reasons['source_url_not_public_http']+=1; continue
        verification=clean(q.get('Mức xác minh'))
        if not verification.startswith(('A','B')): reasons['source_below_B']+=1; continue
        risk=clean(q.get('Rủi ro giáo điều / directive')).lower()
        if risk.startswith(('high','cao')): reasons['directive_risk_high']+=1; continue
        text=clean(q.get('Bản dịch / bản làm việc tiếng Việt'))
        if not 15 <= len(text) <= 600: reasons['text_length_outside_gate']+=1; continue
        if normalized_text(text) in current_text: reasons['exact_text_already_canonical']+=1; continue
        preliminary.append((q,n))
    clusters=defaultdict(list); unclustered=[]
    for q,n in preliminary:
        cluster=clean(n.get('Exact duplicate cluster ID'))
        (clusters[cluster] if cluster else unclustered).append((q,n))
    selected=list(unclustered)
    nature_rank={'direct_quote':0,'translated_quote':1,'proverb':2,'passage':3,'principle':4,'empirical_finding':5,'research_paraphrase':6,'selflo_synthesis':7}
    for members in clusters.values():
        members.sort(key=lambda pair:(0 if clean(pair[0].get('Mức xác minh')).startswith('A') else 1,nature_rank.get(clean(pair[1].get('content_nature canonical')),9),int(pair[1].get('Source row') or 999999)))
        selected.append(members[0]); reasons['exact_duplicate_alternates_held']+=len(members)-1
    selected.sort(key=lambda pair:int(pair[1].get('Source row') or 999999))
    catalog_path=root/'quote-research/review-pipeline/review-candidates.json'; catalog=read(catalog_path); existing={x['id']:x for x in catalog['candidates']}
    chosen=[]
    for q,n in selected:
        original_id=clean(q.get('Quote ID')); ident=stable_id(original_id)
        candidate={'id':ident,'research_id_original':original_id,'text_vi':clean(q.get('Bản dịch / bản làm việc tiếng Việt')),'original_text':clean(q.get('Nguyên văn')),'author':clean(n.get('Author canonical') or q.get('Tác giả / Attribution')),'work':clean(n.get('Source/work canonical') or q.get('Tên nguồn / tác phẩm')),'source_url':uri(q.get('Source URL')),'content_nature':clean(n.get('content_nature canonical') or q.get('Hình thức nội dung (Content Form)')),'life_area':clean(q.get('Lĩnh vực cuộc sống (Life Area)')),'human_experience':clean(n.get('Human experience canonical') or q.get('Trải nghiệm con người (Human Experience)')),'theme':clean(n.get('Theme canonical') or q.get('Chủ đề chính (Theme)')),'verification':clean(q.get('Mức xác minh')),'rights_research_note':clean(q.get('Quyền sử dụng') or q.get('Ghi chú copyright / usage')),'source_row':n.get('Source row'),'duplicate_cluster_ids':clean(n.get('Exact duplicate cluster ID')),'near_duplicate_cluster_ids':clean(n.get('Near-duplicate cluster ID')),'canonical_collision':False,'story_id':None}
        chosen.append(candidate)
        if ident in existing: existing[ident].update(candidate)
        else: catalog['candidates'].append(candidate); existing[ident]=candidate
    now=dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
    ledger={'schema_version':'selflo.content-review-ledger.v1','library_id':'selflo.perspective.candidate.vi','decisions':[{'quote_id':x['id'],'decision':'approved','note':'Full-corpus AI review passed: source A/B, Selflo Fit High, no canonical exact duplicate.'} for x in chosen]}
    audit={'schema_version':'selflo.full-corpus-authoring-audit.v1','generated_at':now,'source_workbook':workbook.name,'research_total':len(raw),'canonical_before':len(current),'selected_for_authoring':len(chosen),'excluded_or_held':len(raw)-len(chosen),'gate':{'selflo_fit':'High','verification':'A or B','source_url':'required','exact_duplicate_policy':'one representative unless already canonical','directive_risk':'exclude High','release_change':False},'reason_counts':dict(sorted(reasons.items())),'selected_ids':[x['id'] for x in chosen],'research_to_canonical_id':{x['research_id_original']:x['id'] for x in chosen if x['research_id_original']!=x['id']}}
    print(json.dumps({'research_total':len(raw),'canonical_before':len(current),'selected_for_authoring':len(chosen),'reason_counts':audit['reason_counts']},ensure_ascii=False,indent=2))
    if args.write:
        write(catalog_path,catalog); write(args.ledger,ledger); write(root/'quote-research/review-pipeline/full-corpus-authoring-audit.json',audit)


if __name__=='__main__': main()
